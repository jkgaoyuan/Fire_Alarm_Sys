"""
设备历史轨迹与导出（3.3 B-18 / FR-018）测试

覆盖：时间区间边界与升序、默认窗口与 90 天上限校验、xlsx/csv 导出行数一致。
"""

import io
from datetime import datetime, timedelta

import pytest
import pytest_asyncio
from openpyxl import load_workbook

from app.models.device import DeviceStatusLog
from tests.device_helpers import (
    ALL_DEVICE_PERMS,
    auth_headers,
    create_device_type,
    create_device_user,
    create_org,
    device_payload,
)

BASE = datetime(2026, 1, 1, 0, 0, 0)


@pytest_asyncio.fixture
async def trajectory_env(db_session):
    org = await create_org(db_session, "轨迹大楼")
    device_type = await create_device_type(db_session)
    chief = await create_device_user(
        db_session, username="traj_chief", perm_codes=ALL_DEVICE_PERMS, data_scope="all"
    )
    return {
        "org": org,
        "type": device_type,
        "user": chief,
        "payload": device_payload(device_type.id, org.id, "DEV-TRJ-001"),
    }


async def _device_with_logs(client, env, db, times: list[datetime]) -> int:
    """建档设备 + 指定时刻的状态变更日志（固定时刻，区间断言才可复现）"""
    resp = await client.post(
        "/api/v1/devices", headers=auth_headers(env["user"]), json=env["payload"]
    )
    assert resp.status_code == 200, resp.text
    device_id = resp.json()["data"]["id"]

    for index, moment in enumerate(times):
        db.add(
            DeviceStatusLog(
                device_id=device_id,
                old_status="normal" if index % 2 == 0 else "fault",
                new_status="fault" if index % 2 == 0 else "normal",
                reason=f"第{index + 1}次变更",
                created_at=moment,
            )
        )
    await db.commit()
    return device_id


@pytest.mark.asyncio
async def test_trajectory_window_boundaries(client, trajectory_env, db_session):
    """闭区间 [start, end] 命中、区间外排除、结果按时间升序"""
    env = trajectory_env
    times = [BASE, BASE + timedelta(days=2), BASE + timedelta(days=4), BASE + timedelta(days=9)]
    device_id = await _device_with_logs(client, env, db_session, times)
    headers = auth_headers(env["user"])

    data = (
        await client.get(
            f"/api/v1/devices/{device_id}/trajectory",
            headers=headers,
            params={
                "start": BASE.isoformat(),
                "end": (BASE + timedelta(days=4)).isoformat(),
            },
        )
    ).json()["data"]
    assert data["total"] == 3
    assert [i["reason"] for i in data["items"]] == ["第1次变更", "第2次变更", "第3次变更"]
    item_times = [i["time"] for i in data["items"]]
    assert item_times == sorted(item_times)
    assert data["items"][0]["status_label"] == "故障"
    assert data["items"][0]["old_status"] == "normal"

    # 建档日志落在 now：默认近 7 天窗口只有它，Jan 区间内则没有它
    default = (
        await client.get(f"/api/v1/devices/{device_id}/trajectory", headers=headers)
    ).json()["data"]
    assert default["total"] == 1
    assert default["items"][0]["old_status"] is None


@pytest.mark.asyncio
async def test_trajectory_rejects_invalid_window(client, trajectory_env, db_session):
    """跨度超 90 天或起止倒置返回 code=400；设备不存在返回 404"""
    env = trajectory_env
    device_id = await _device_with_logs(client, env, db_session, [BASE])
    headers = auth_headers(env["user"])

    too_wide = (
        await client.get(
            f"/api/v1/devices/{device_id}/trajectory",
            headers=headers,
            params={
                "start": BASE.isoformat(),
                "end": (BASE + timedelta(days=91)).isoformat(),
            },
        )
    ).json()
    assert too_wide["code"] == 400 and "90" in too_wide["message"]

    inverted = (
        await client.get(
            f"/api/v1/devices/{device_id}/trajectory",
            headers=headers,
            params={
                "start": (BASE + timedelta(days=1)).isoformat(),
                "end": BASE.isoformat(),
            },
        )
    ).json()
    assert inverted["code"] == 400

    missing = await client.get("/api/v1/devices/424242/trajectory", headers=headers)
    assert missing.json()["code"] == 404


@pytest.mark.asyncio
async def test_trajectory_export_matches_query(client, trajectory_env, db_session):
    """xlsx / csv 导出行数与查询结果一致，表头与状态列可读"""
    env = trajectory_env
    times = [BASE + timedelta(hours=index) for index in range(1, 6)]
    device_id = await _device_with_logs(client, env, db_session, times)
    headers = auth_headers(env["user"])
    params = {
        "start": (BASE - timedelta(minutes=1)).isoformat(),
        "end": (BASE + timedelta(days=1)).isoformat(),
    }

    total = (
        await client.get(
            f"/api/v1/devices/{device_id}/trajectory", headers=headers, params=params
        )
    ).json()["data"]["total"]
    assert total == 5

    xlsx = await client.get(
        f"/api/v1/devices/{device_id}/trajectory/export",
        headers=headers,
        params={**params, "format": "xlsx"},
    )
    assert xlsx.status_code == 200
    assert "spreadsheetml.sheet" in xlsx.headers["content-type"]
    assert "DEV-TRJ-001_trajectory_" in xlsx.headers["content-disposition"]
    rows = list(load_workbook(io.BytesIO(xlsx.content)).active.iter_rows(values_only=True))
    assert rows[0] == ("时间", "原状态", "新状态", "变更原因", "操作人")
    assert len(rows) - 1 == total
    assert rows[1][3] == "第1次变更" and rows[1][4] == "系统"

    csv_resp = await client.get(
        f"/api/v1/devices/{device_id}/trajectory/export",
        headers=headers,
        params={**params, "format": "csv"},
    )
    assert csv_resp.status_code == 200
    lines = [line for line in csv_resp.content.decode("utf-8-sig").splitlines() if line]
    assert len(lines) - 1 == total
    assert lines[1].split(",")[3] == "第1次变更"
