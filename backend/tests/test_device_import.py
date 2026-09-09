"""
设备档案批量导入（FR-009 / B-10）测试

覆盖：模板下载、正常导入、编码重复、区域不存在、部分导入、扩展属性校验失败明细、
失败率超阈值整体回滚。
"""

import io

import pytest
import pytest_asyncio
from openpyxl import Workbook, load_workbook

from app.services.device_import_service import FIXED_COLUMNS, SHEET_NAME
from tests.device_helpers import (
    auth_headers,
    create_device_type,
    create_device_user,
    create_org,
    device_payload,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_xlsx(rows: list[dict]) -> bytes:
    """按导入模板的表头顺序生成一个工作簿"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append([header for header, _ in FIXED_COLUMNS])
    for row in rows:
        ws.append([row.get(field) for _, field in FIXED_COLUMNS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def valid_row(code: str, type_code: str, org_id: int, **overrides) -> dict:
    row = {
        "device_code": code,
        "device_name": f"设备{code}",
        "type_code": type_code,
        "org_id": org_id,
        "manufacturer": "霍尼韦尔",
        "model": "XLS-PS",
        "brand": "Honeywell",
        "spec": "光电型",
        "install_date": "2025-03-15",
        "warranty_expire_date": "2028-03-15",
        "maintain_cycle": 90,
        "status": "normal",
        "map_x": 120.5,
        "map_y": 340.2,
        "remark": None,
        "attributes": '{"sensitivity": "高", "detection_area": 60}',
    }
    row.update(overrides)
    return row


@pytest_asyncio.fixture
async def import_env(db_session):
    org = await create_org(db_session, "总部大楼")
    await create_device_type(db_session, "smoke_detector", "烟感探测器")
    await create_device_type(
        db_session, "hydrant", "消火栓", {"outlet_count": {"label": "水带数量", "type": "number"}}
    )
    chief = await create_device_user(
        db_session,
        username="chief",
        perm_codes=["device:create", "device:view", "device:delete"],
        data_scope="all",
    )
    return {"org": org, "user": chief}


async def _upload(client, env, content: bytes, headers=None):
    """上传导入。headers 需在导入前取好——整批回滚会让会话对象过期"""
    return await client.post(
        "/api/v1/devices/import",
        headers=headers or auth_headers(env["user"]),
        files={"file": ("devices.xlsx", content, XLSX_MIME)},
    )


@pytest.mark.asyncio
async def test_download_template(client, import_env):
    """模板含全部列头与 8 种类型说明，可直接被导入接口识别"""
    resp = await client.get(
        "/api/v1/devices/import/template", headers=auth_headers(import_env["user"])
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(XLSX_MIME)
    assert "attachment" in resp.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb[SHEET_NAME]
    headers = [cell.value for cell in ws[1]]
    assert headers == [header for header, _ in FIXED_COLUMNS]
    assert wb["填写说明"]["A2"].value.startswith("1.")


@pytest.mark.asyncio
async def test_import_all_valid_rows(client, import_env):
    """全部合法时逐行建档，类型与区域解析正确"""
    env = import_env
    org_id = env["org"].id
    content = build_xlsx(
        [
            valid_row("DEV-IMP-001", "smoke_detector", org_id),
            valid_row("DEV-IMP-002", "smoke_detector", org_id),
            valid_row("DEV-IMP-003", "hydrant", org_id, attributes='{"outlet_count": 2}'),
        ]
    )

    resp = await _upload(client, env, content)
    assert resp.status_code == 200
    data = resp.json()["data"]
    assert (data["total"], data["success"], data["failed"]) == (3, 3, 0)
    assert data["rolled_back"] is False

    listed = (
        await client.get("/api/v1/devices?page_size=100", headers=auth_headers(env["user"]))
    ).json()["data"]
    assert listed["total"] == 3
    hydrant = next(d for d in listed["items"] if d["device_code"] == "DEV-IMP-003")
    assert hydrant["type_name"] == "消火栓"
    assert hydrant["attributes"] == {"outlet_count": 2}
    assert hydrant["org_name"] == "总部大楼"
    assert hydrant["creator_name"] == "chief"


@pytest.mark.asyncio
async def test_import_writes_archive_status_log(client, import_env):
    """导入建档同样写入状态留痕，否则导入设备的历史时间轴为空（FR-011）"""
    env = import_env
    content = build_xlsx([valid_row("DEV-IMP-H1", "smoke_detector", env["org"].id)])
    await _upload(client, env, content)

    items = (
        await client.get("/api/v1/devices?page_size=10", headers=auth_headers(env["user"]))
    ).json()["data"]["items"]
    device_id = next(item["id"] for item in items if item["device_code"] == "DEV-IMP-H1")

    history = (
        await client.get(
            f"/api/v1/devices/{device_id}/history", headers=auth_headers(env["user"])
        )
    ).json()["data"]
    assert history["total"] == 1
    assert history["items"][0]["title"] == "建档：正常"
    assert history["items"][0]["detail"] == "批量导入建档"


@pytest.mark.asyncio
async def test_import_duplicate_code(client, import_env, db_session):
    """与库中已有编码冲突的行失败并标注行号，其余行仍写入"""
    env = import_env
    org_id = env["org"].id
    payload = device_payload(1, org_id, "DEV-EXISTS")
    payload["attributes"] = {}
    await client.post("/api/v1/devices", headers=auth_headers(env["user"]), json=payload)

    content = build_xlsx(
        [
            valid_row("DEV-NEW-001", "smoke_detector", org_id),
            valid_row("DEV-EXISTS", "smoke_detector", org_id),
        ]
    )
    resp = await _upload(client, env, content)
    data = resp.json()["data"]
    assert (data["total"], data["success"], data["failed"]) == (2, 1, 1)
    assert data["failures"] == [
        {"row": 3, "device_code": "DEV-EXISTS", "reason": "设备编码已存在: DEV-EXISTS"}
    ]


@pytest.mark.asyncio
async def test_import_duplicate_deleted_code_reports_restore_entry(
    client, import_env, db_session
):
    """导入时撞到已逻辑删除档案，错误文案应提示恢复入口"""
    env = import_env
    org_id = env["org"].id
    headers = auth_headers(env["user"])
    payload = device_payload(1, org_id, "DEV-EXISTS-DEL")
    payload["attributes"] = {}
    created = (
        await client.post("/api/v1/devices", headers=headers, json=payload)
    ).json()["data"]
    await client.delete(f"/api/v1/devices/{created['id']}", headers=headers)

    content = build_xlsx(
        [
            valid_row("DEV-NEW-002", "smoke_detector", org_id),
            valid_row("DEV-EXISTS-DEL", "smoke_detector", org_id),
        ]
    )
    resp = await _upload(client, env, content)
    data = resp.json()["data"]
    assert (data["total"], data["success"], data["failed"]) == (2, 1, 1)
    failure = data["failures"][0]
    assert failure["row"] == 3
    assert failure["device_code"] == "DEV-EXISTS-DEL"
    assert "设备编码已被已删除档案占用" in failure["reason"]
    assert f"/api/v1/devices/{created['id']}/restore" in failure["reason"]


@pytest.mark.asyncio
async def test_import_unknown_org_and_type(client, import_env):
    """区域不存在与类型编码非法都要逐行报错"""
    env = import_env
    content = build_xlsx(
        [
            valid_row("DEV-BAD-001", "smoke_detector", 999999),
            valid_row("DEV-BAD-002", "no_such_type", env["org"].id),
        ]
    )
    resp = await _upload(client, env, content)
    data = resp.json()["data"]
    assert data["success"] == 0
    assert [(f["row"], f["reason"]) for f in data["failures"]] == [
        (2, "区域不存在: id=999999"),
        (3, "设备类型不存在: no_such_type"),
    ]


@pytest.mark.asyncio
async def test_import_attribute_validation_detail(client, import_env):
    """扩展属性非法值要在失败明细里给出字段级原因"""
    env = import_env
    content = build_xlsx(
        [
            valid_row("DEV-ATTR-001", "smoke_detector", env["org"].id),
            valid_row("DEV-ATTR-002", "smoke_detector", env["org"].id),
            valid_row("DEV-ATTR-003", "smoke_detector", env["org"].id),
            valid_row("DEV-ATTR-004", "smoke_detector", env["org"].id),
            valid_row("DEV-ATTR-005", "smoke_detector", env["org"].id,
                      attributes='{"sensitivity": "极高"}'),
            valid_row("DEV-ATTR-006", "smoke_detector", env["org"].id,
                      attributes="not-a-json"),
            valid_row("DEV-ATTR-007", "smoke_detector", env["org"].id,
                      attributes='{"unknown_field": 1}'),
        ]
    )
    resp = await _upload(client, env, content)
    data = resp.json()["data"]
    assert (data["total"], data["success"], data["failed"]) == (7, 4, 3)

    reasons = {f["row"]: f["reason"] for f in data["failures"]}
    assert "灵敏度 取值必须是 ['高', '中', '低'] 之一" in reasons[6]
    assert reasons[7] == "扩展属性(JSON) 格式错误"
    assert "未知扩展属性: unknown_field" in reasons[8]


@pytest.mark.asyncio
async def test_import_rolls_back_when_failure_rate_too_high(client, import_env):
    """失败率超过 50% 时整批回滚，一行也不落库"""
    env = import_env
    content = build_xlsx(
        [
            valid_row("DEV-RB-001", "smoke_detector", env["org"].id),
            valid_row("DEV-RB-002", "smoke_detector", 999999),
            valid_row("DEV-RB-003", "smoke_detector", 999999),
            valid_row("DEV-RB-004", "smoke_detector", 999999),
        ]
    )
    headers = auth_headers(env["user"])
    resp = await _upload(client, env, content, headers)
    data = resp.json()["data"]
    assert data["rolled_back"] is True
    assert data["success"] == 0
    assert data["failed"] == 4
    assert {f["reason"] for f in data["failures"]} == {
        "区域不存在: id=999999",
        "失败率超过阈值，整批回滚",
    }

    listed = (await client.get("/api/v1/devices", headers=headers)).json()["data"]
    assert listed["total"] == 0


@pytest.mark.asyncio
async def test_import_rejects_bad_file(client, import_env):
    """缺少必需列 / 非 Excel 内容都返回 400 而不是 500"""
    env = import_env

    wb = Workbook()
    wb.active.append(["无关列"])
    buf = io.BytesIO()
    wb.save(buf)
    resp = await _upload(client, env, buf.getvalue())
    assert resp.status_code == 400
    assert "缺少必需列" in resp.json()["message"]

    resp = await _upload(client, env, b"plain text, not a workbook")
    assert resp.status_code == 400
    assert "无法解析 Excel" in resp.json()["message"]
