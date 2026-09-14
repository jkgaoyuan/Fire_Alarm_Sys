"""
报表导出服务（3.9 FR-052）
========================================
异步导出框架，支持：
- Excel/CSV：流式写入 openpyxl write_only=True 模式
- Word：python-docx，演练报告生成，图片≤30 张（OQ-3）
- BackgroundTasks：首版不实现任务恢复（OQ-6），失败标记为 failed 供重试
"""

import io
import os
import uuid
from datetime import datetime
from typing import Optional

from fastapi import BackgroundTasks
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.report_export import ReportExportTask


EXPORT_TASK_TYPES = (
    "alarm_trend",
    "device_status",
    "fault_top10",
    "inspection",
    "drill_report",
)


def _generate_task_no() -> str:
    """生成唯一任务编号"""
    return f"EXP_{datetime.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex[:8].upper()}"


async def create_export_task(
    db: AsyncSession,
    task_type: str,
    params: dict,
    created_by: int,
) -> ReportExportTask:
    """创建导出任务记录"""
    task = ReportExportTask(
        task_no=_generate_task_no(),
        task_type=task_type,
        status="pending",
        params=params,
        created_by=created_by,
    )
    db.add(task)
    await db.flush()
    await db.refresh(task)
    return task


async def execute_export_sync(
    db: AsyncSession,
    task: ReportExportTask,
    data: list,
    file_format: str = "xlsx",
) -> ReportExportTask:
    """同步执行导出（≤1 万行数据）

    直接在当前请求中完成文件生成和状态更新。
    """
    task.status = "running"
    task.total_rows = len(data)
    await db.flush()

    try:
        if task.task_type == "drill_report" or file_format == "docx":
            file_path, file_name = await _generate_word(db, task, data)
        else:
            file_path, file_name = await _generate_excel(db, task, data)

        task.status = "completed"
        task.file_path = file_path
        task.file_name = file_name
        task.completed_at = datetime.utcnow()
    except Exception as e:
        task.status = "failed"
        task.error_message = str(e)

    # ⚠️ 必须 commit（本仓库约定：由 service 层提交，见其他 service）。
    # `get_db` 不自动提交，且在 finally 里 close() 会话 —— 只 flush 的后果是
    # **接口返回 task_id、文件也已写进磁盘，而库里 0 行**（事务被回滚），
    # 导出中心因此永远查不到任何任务。实测：接口返回 task_id=19，库里 0 行。
    await db.commit()

    return task


async def execute_export_async(
    background_tasks: BackgroundTasks,
    db: AsyncSession,
    task: ReportExportTask,
    data: list,
    file_format: str = "xlsx",
) -> ReportExportTask:
    """异步执行导出（>1 万行数据，通过 BackgroundTasks）"""
    task.status = "running"
    task.total_rows = len(data)
    await db.flush()

    # 将实际导出逻辑放入后台任务
    background_tasks.add_task(
        _background_export, task.id, data, file_format
    )

    # ⚠️ 同上：必须在返回前提交。`_background_export` 用的是**独立 session**，
    # 靠 `db.get(ReportExportTask, task_id)` 取回本行；不提交的话它拿到 None
    # 直接 return，任务永久停在 running 而用户看不到任何错误。
    await db.commit()

    return task


async def _background_export(task_id: int, data: list, file_format: str):
    """后台导出任务（独立 session）"""
    # ⚠️ 符号名是 AsyncSessionLocal。这里此前写的是 `async_session_maker`——
    # 该名字在 app/db/session.py 中并不存在，导致后台导出**第一行就 ImportError**，
    # 异常被后台任务机制吞掉、任务永久停在 running（>1 万行导出 100% 死路）。
    from app.db.session import AsyncSessionLocal

    async with AsyncSessionLocal() as db:
        task = await db.get(ReportExportTask, task_id)
        if not task:
            return

        try:
            if task.task_type == "drill_report" or file_format == "docx":
                file_path, file_name = await _generate_word(db, task, data)
            else:
                file_path, file_name = await _generate_excel(db, task, data)

            task.status = "completed"
            task.file_path = file_path
            task.file_name = file_name
            task.completed_at = datetime.utcnow()
        except Exception as e:
            task.status = "failed"
            task.error_message = str(e)

        await db.flush()
        await db.commit()


async def _generate_excel(
    db: AsyncSession,
    task: ReportExportTask,
    data: list,
) -> tuple[str, str]:
    """生成 Excel 文件"""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("数据")

    # 写入表头
    if data and isinstance(data[0], dict):
        headers = list(data[0].keys())
        ws.append(headers)
        for row in data:
            ws.append([row.get(h, "") for h in headers])
    else:
        ws.append(["数据"])
        for item in data:
            ws.append([str(item)])

    # 保存文件
    save_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "storage", "export")
    os.makedirs(save_dir, exist_ok=True)

    name_map = {
        "alarm_trend": "报警趋势表",
        "device_status": "设备状态统计表",
        "fault_top10": "故障TOP10表",
        "inspection": "巡检完成率表",
    }
    prefix = name_map.get(task.task_type, "统计表")
    date_str = datetime.now().strftime("%Y%m%d")
    file_name = f"{prefix}_{date_str}.xlsx"

    # ⚠️ 磁盘路径必须带 task_no（唯一约束），不能直接用 file_name：
    # file_name 只含「类型 + 日期」，**同类型同天的两个任务必然同名**，
    # 后者覆盖前者；而两条 DB 记录都指向同一路径，于是下载历史任务会
    # 静默拿到后一次的内容（实测：任务 A 记着 total_rows=2，下载到的文件
    # 却只有 1 行 —— 那是任务 B 的数据）。
    # 用 task_no 拼路径、file_name 留给界面与下载名，两个关注点分开。
    file_path = os.path.join(save_dir, f"{task.task_no}_{file_name}")

    wb.save(file_path)
    return file_path, file_name


async def _generate_word(
    db: AsyncSession,
    task: ReportExportTask,
    data: list,
) -> tuple[str, str]:
    """生成 Word 演练报告"""
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # 标题
    title = doc.add_heading("消防演练评估报告", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    params = task.params or {}

    # 基本信息
    doc.add_heading("一、演练基本信息", level=1)
    doc.add_paragraph(f"演练名称：{params.get('drill_name', 'N/A')}")
    doc.add_paragraph(f"演练类型：{params.get('drill_type', 'N/A')}")
    doc.add_paragraph(f"演练时间：{params.get('actual_start_at', 'N/A')}")

    # 评估项打分
    doc.add_heading("二、评估项打分", level=1)
    items = data if isinstance(data, list) else []
    if items and isinstance(items[0], dict):
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "评估项"
        hdr[1].text = "满分"
        hdr[2].text = "得分"
        hdr[3].text = "评语"
        for item in items[:50]:
            row = table.add_row().cells
            row[0].text = str(item.get("label", "N/A"))
            row[1].text = str(item.get("max_score", 10))
            row[2].text = str(item.get("score", 0))
            row[3].text = str(item.get("comment", ""))

    # 问题与改进
    doc.add_heading("三、存在问题与改进措施", level=1)
    if params.get("problems"):
        doc.add_paragraph(f"存在问题:\n{params['problems']}")
    if params.get("improvements"):
        doc.add_paragraph(f"改进措施:\n{params['improvements']}")

    # 总体评估
    doc.add_heading("四、总体评估摘要", level=1)
    if params.get("evaluation_summary"):
        doc.add_paragraph(params["evaluation_summary"])

    # 照片附件（OQ-3: 最多 30 张）
    photos = params.get("photos", [])
    if photos:
        doc.add_heading("五、现场照片", level=1)
        for photo in photos[:30]:
            url = photo.get("url", "")
            caption = photo.get("caption", "")
            if url and os.path.exists(url):
                try:
                    doc.add_picture(url, height=Inches(3))
                    doc.add_paragraph(caption).alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    doc.add_paragraph(f"[图片加载失败] {caption}")
            else:
                doc.add_paragraph(f"[图片占位] {caption}")

        total_photos = len(photos)
        if total_photos > 30:
            doc.add_paragraph(f"\n共 {total_photos} 张照片，本报告仅展示前 30 张。")

    # 保存
    save_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "storage", "export")
    os.makedirs(save_dir, exist_ok=True)

    drill_name = params.get("drill_name", "演练").replace("/", "_")[:30]
    date_str = datetime.now().strftime("%Y%m%d")
    file_name = f"演练报告_{drill_name}_{date_str}.docx"

    # 同上：同一天对同一场演练重复导出会撞名并互相覆盖
    file_path = os.path.join(save_dir, f"{task.task_no}_{file_name}")

    doc.save(file_path)
    return file_path, file_name
