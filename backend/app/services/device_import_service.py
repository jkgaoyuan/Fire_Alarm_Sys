"""
设备档案批量导入（3.2 B-10）

Excel 解析与校验（编码唯一 / 区域存在 / 类型有效 / 扩展属性 JSON Schema）→
逐行写入（单行失败不中断）→ 失败率超过阈值时整体回滚。
"""

import io
import json
from datetime import date
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AuthError
from app.crud.device import device_crud
from app.models.device import DEVICE_STATUSES, Device
from app.models.device_type import DeviceType
from app.models.organization import Organization
from app.models.user import User
from app.services.device_service import validate_attributes, write_status_log

# 失败率超过该阈值时整体回滚（计划后端技术要点 3）
ROLLBACK_FAILURE_RATE = 0.5

SHEET_NAME = "设备档案"

# 固定列：(表头, 字段名)。扩展属性以 JSON 文本承载在最后一列。
FIXED_COLUMNS: list[tuple[str, str]] = [
    ("设备编码*", "device_code"),
    ("设备名称*", "device_name"),
    ("类型编码", "type_code"),
    ("区域ID", "org_id"),
    ("厂商", "manufacturer"),
    ("型号", "model"),
    ("品牌", "brand"),
    ("规格", "spec"),
    ("安装日期", "install_date"),
    ("质保到期日", "warranty_expire_date"),
    ("维护周期(天)", "maintain_cycle"),
    ("状态", "status"),
    ("X坐标", "map_x"),
    ("Y坐标", "map_y"),
    ("备注", "remark"),
    ("扩展属性(JSON)", "attributes"),
]

REQUIRED_FIELDS = ("device_code", "device_name")
STRING_FIELDS = ("manufacturer", "model", "brand", "spec", "remark")
EXAMPLE_ATTRIBUTES = '{"sensitivity": "高", "detection_area": 60}'


class RowError(Exception):
    """单行校验失败：记入失败明细后继续处理下一行"""


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _to_date(value: Any) -> date | None:
    """openpyxl 会把日期单元格解析为 datetime，需与文本格式一并支持"""
    if value is None or _text(value) is None:
        return None
    if isinstance(value, date) and not isinstance(value, str):
        return value.date() if hasattr(value, "date") else value
    text = str(value).strip().replace("/", "-").split(" ")[0]
    parts = text.split("-")
    if len(parts) >= 3:
        try:
            return date(int(parts[0]), int(parts[1]), int(parts[2]))
        except ValueError:
            pass
    raise RowError(f"日期格式无法解析: {value}")


def _integer(value: Any, label: str) -> int | None:
    text = _text(value)
    if text is None:
        return None
    try:
        return int(float(text))
    except ValueError:
        raise RowError(f"{label} 必须为整数: {text}") from None


def _float(value: Any, label: str) -> float | None:
    text = _text(value)
    if text is None:
        return None
    try:
        return float(text)
    except ValueError:
        raise RowError(f"{label} 必须为数字: {text}") from None


# ==================== 模板生成 ====================


def build_template(device_types: list[DeviceType]) -> bytes:
    """生成导入模板：表头 + 示例行 + 填写说明页"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    headers = [header for header, _ in FIXED_COLUMNS]
    ws.append(headers)
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(idx)].width = max(
            14, len(header) * 2 + 4
        )

    ws.append(
        [
            "DEV-SMK-001",
            "1F大厅烟感A01",
            device_types[0].type_code if device_types else "smoke_detector",
            1,
            "霍尼韦尔",
            "XLS-PS",
            "Honeywell",
            "光电型",
            "2025-03-15",
            "2028-03-15",
            90,
            "normal",
            120.5,
            340.2,
            "示例行，导入前请删除",
            EXAMPLE_ATTRIBUTES,
        ]
    )

    type_hint = "、".join(f"{t.type_code}（{t.type_name}）" for t in device_types)
    notes = wb.create_sheet("填写说明")
    notes.column_dimensions["A"].width = 110
    for line in [
        "填写说明：",
        "1. 「设备编码」「设备名称」为必填，编码全库唯一。",
        f"2. 「类型编码」取值：{type_hint or '（设备类型表为空，请先初始化数据）'}",
        "3. 「区域ID」为组织架构节点的 id，可在系统组织架构中查询。",
        "4. 日期格式 YYYY-MM-DD 或 YYYY/MM/DD。",
        f"5. 「状态」取值：{', '.join(DEVICE_STATUSES)}；留空按 normal 处理。",
        f"6. 「扩展属性(JSON)」须属于该设备类型的属性模板，例如 {EXAMPLE_ATTRIBUTES}。",
        f"7. 单行失败不影响其他行写入；失败率超过 {int(ROLLBACK_FAILURE_RATE * 100)}% 时整批回滚。",
    ]:
        notes.append([line])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==================== Excel 解析 ====================


def read_rows(payload: bytes) -> list[tuple[int, dict[str, Any]]]:
    """读取工作表，返回 [(Excel 行号, 字段字典)]，跳过全空行"""
    try:
        wb = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
    except Exception:
        raise AuthError(400, "无法解析 Excel 文件，请使用系统模板") from None

    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows or not any(str(h).strip() for h in rows[0] if h is not None):
        raise AuthError(400, "Excel 表头缺失或内容为空")

    headers = {
        str(header).strip(): idx
        for idx, header in enumerate(rows[0])
        if header is not None
    }
    missing = [
        header
        for header, field in FIXED_COLUMNS
        if field in REQUIRED_FIELDS and header not in headers
    ]
    if missing:
        raise AuthError(400, f"Excel 缺少必需列: {', '.join(missing)}")

    parsed: list[tuple[int, dict[str, Any]]] = []
    for offset, values in enumerate(rows[1:], start=2):
        if values is None or all(
            v is None or str(v).strip() == "" for v in values
        ):
            continue
        row: dict[str, Any] = {}
        for header, field in FIXED_COLUMNS:
            idx = headers.get(header)
            if idx is not None and idx < len(values):
                row[field] = values[idx]
        parsed.append((offset, row))
    return parsed


async def _load_lookup(db: AsyncSession) -> tuple[dict[str, DeviceType], set[int]]:
    """一次取出类型映射与区域 id 集合，避免逐行查库"""
    type_rows = (await db.execute(select(DeviceType))).scalars().all()
    org_rows = (await db.execute(select(Organization.id))).scalars().all()
    return {t.type_code: t for t in type_rows}, set(org_rows)


def parse_row(
    row: dict[str, Any],
    type_map: dict[str, DeviceType],
    org_ids: set[int],
    taken_codes: set[str],
) -> Device:
    """一行 Excel → 待插入的 Device 实例；任一校验不通过抛 RowError"""
    device_code = _text(row.get("device_code"))
    if not device_code:
        raise RowError("设备编码不能为空")
    if device_code in taken_codes:
        raise RowError(f"设备编码已存在: {device_code}")

    device_name = _text(row.get("device_name"))
    if not device_name:
        raise RowError("设备名称不能为空")

    device_type: DeviceType | None = None
    type_code = _text(row.get("type_code"))
    if type_code:
        device_type = type_map.get(type_code)
        if device_type is None:
            raise RowError(f"设备类型不存在: {type_code}")

    org_id: int | None = None
    raw_org = _text(row.get("org_id"))
    if raw_org:
        org_id = _integer(raw_org, "区域ID")
        if org_id not in org_ids:
            raise RowError(f"区域不存在: id={org_id}")

    status = _text(row.get("status")) or "normal"
    if status not in DEVICE_STATUSES:
        raise RowError(f"状态非法: {status}")

    attributes: dict[str, Any] = {}
    raw_attrs = _text(row.get("attributes"))
    if raw_attrs:
        try:
            parsed = json.loads(raw_attrs)
        except json.JSONDecodeError:
            raise RowError("扩展属性(JSON) 格式错误") from None
        if not isinstance(parsed, dict):
            raise RowError("扩展属性(JSON) 必须是对象")
        attributes = parsed

    schema = (device_type.attribute_schema or {}) if device_type else {}
    errors = validate_attributes(schema, attributes)
    if errors:
        raise RowError("扩展属性校验失败: " + "；".join(errors))

    device = Device(
        device_code=device_code,
        device_name=device_name,
        type_id=device_type.id if device_type else None,
        org_id=org_id,
        status=status,
        attributes=attributes,
        maintain_cycle=_integer(row.get("maintain_cycle"), "维护周期"),
        map_x=_float(row.get("map_x"), "X坐标"),
        map_y=_float(row.get("map_y"), "Y坐标"),
        install_date=_to_date(row.get("install_date")),
        warranty_expire_date=_to_date(row.get("warranty_expire_date")),
    )
    for field in STRING_FIELDS:
        setattr(device, field, _text(row.get(field)))
    device.remark = _text(row.get("remark"))
    return device


# ==================== 导入主流程 ====================


async def import_devices(
    db: AsyncSession, payload: bytes, user: User
) -> dict[str, Any]:
    """
    批量导入设备，返回 {total, success, failed, failures, rolled_back, message}。

    先逐行构建实例（不入库），全部解析完成后再统一 add + commit；
    失败率超阈值则直接 rollback 且不写入任何一行，因此不会留下脏数据。
    """
    rows = read_rows(payload)
    if not rows:
        raise AuthError(400, "Excel 中没有可导入的数据行")

    type_map, org_ids = await _load_lookup(db)
    codes = [code for code in (_text(r.get("device_code")) for _, r in rows) if code]
    taken_codes = await device_crud.get_codes_in_use(db, codes)

    failures: list[dict[str, Any]] = []
    pending: list[tuple[int, Device]] = []

    for row_no, row in rows:
        code = _text(row.get("device_code"))
        try:
            device = parse_row(row, type_map, org_ids, taken_codes)
        except RowError as exc:
            failures.append({"row": row_no, "device_code": code, "reason": str(exc)})
            continue
        taken_codes.add(device.device_code)
        device.created_by = user.id
        pending.append((row_no, device))

    total = len(rows)
    failed = total - len(pending)
    if total and failed / total > ROLLBACK_FAILURE_RATE:
        await db.rollback()
        for row_no, device in pending:
            failures.append(
                {
                    "row": row_no,
                    "device_code": device.device_code,
                    "reason": "失败率超过阈值，整批回滚",
                }
            )
        return {
            "total": total,
            "success": 0,
            "failed": total,
            "failures": sorted(failures, key=lambda f: f["row"]),
            "rolled_back": True,
            "message": (
                f"失败 {failed}/{total} 行，超过 "
                f"{int(ROLLBACK_FAILURE_RATE * 100)}%，已整体回滚，请修正后重新导入"
            ),
        }

    for _, device in pending:
        db.add(device)
    await db.flush()
    for _, device in pending:
        await write_status_log(
            db,
            device_id=device.id,
            old_status=None,
            new_status=device.status,
            changed_by=user.id,
            reason="批量导入建档",
        )
    await db.commit()

    return {
        "total": total,
        "success": len(pending),
        "failed": failed,
        "failures": sorted(failures, key=lambda f: f["row"]),
        "rolled_back": False,
        "message": "导入完成",
    }
